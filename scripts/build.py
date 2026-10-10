import io, os, re, sys, json, base64, hashlib, shutil, pathlib, html, datetime
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageOps

ROOT = pathlib.Path(".")
SITE = "https://fulllife.blog"
MODEL = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
# 含這些字的投影片（謄本、使用執照等）完全不處理：文字不送 AI、圖片不公開
SENSITIVE = ["使用執照", "謄本", "測量成果", "登記簿", "所有權狀", "身分證", "契約"]
MAX_CAND = 30
CATS = ["大型廠房", "小型廠房", "土地", "其他"]
DATA = ROOT / "data" / "listings.json"
# 樣式表版本號：style.css 改了就換，避免瀏覽器用舊快取
CSS_V = hashlib.md5((ROOT / "style.css").read_bytes()).hexdigest()[:8] if (ROOT / "style.css").exists() else "1"
JS_V = hashlib.md5((ROOT / "app.js").read_bytes()).hexdigest()[:8] if (ROOT / "app.js").exists() else "1"

PROMPT = """你是工業不動產物件資料整理員。以下是一份物件簡報的文字（已排除謄本、使用執照等頁面）與候選圖片。
請只輸出 JSON，欄位如下：
title：物件標題，例如「仁武區 工業廠房 小坪數好用廠房出售」，不含公司名
type："售" 或 "租"
price：價格，照原文（如 "5,300萬（每坪約 32萬）"、"月租 45.35萬"、"160萬/坪"），不要自行計算
land_ping：土地坪數，只填數字字串；沒有則 ""
build_ping：建物坪數，只填數字字串；沒有則 ""
zoning：使用分區，如 "乙種工業區"、"丁種建築用地"、"住四(50%/300%)"
area：公開地址，只到「縣市＋區/鄉/鎮＋路/街/巷」，縣市必須依原文地址填寫（可能是屏東縣、台南市等，不要預設高雄），不得含門牌號碼、巷弄號、地號、段號
category："大型廠房"（廠房/廠辦/倉庫，建坪約500坪以上）、"小型廠房"（建坪500坪以下的廠房/倉庫）、"土地"（無建物的土地）、"其他"（住宅、店面、辦公、車位等非廠房）擇一
note：公開備註（面寬、深度、路寬、樓高、電力、天車、載重、完工日、結構、是否帶租約等客觀資料，用「、」分隔），沒有則 ""
desc：投資亮點／訴求重點，每點一句，用「；」分隔
photo_ids：候選圖片中屬於「實景照片」的編號（整數陣列），依適合展示的順序排列（外觀、空拍優先），最多 12 個
規則：
- 實景照片＝建物外觀、內部、空拍、周邊街景實拍。排除：地圖、地籍圖、平面圖、配置圖、證件文件、謄本、截圖、logo、橫幅、人像。
- 絕對不得寫入任何欄位：屋主姓名、同行出價、議價空間、「後台內容」等內部備註、電話、公司名稱。
- 原文沒有的欄位填空字串，不要猜測或編造。"""


def walk(shapes):
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(s.shapes)
        else:
            yield s


def slide_lines(slide):
    out = []
    for s in walk(slide.shapes):
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    out.append(t)
    return out


def is_sensitive(lines):
    """標題式的短行含關鍵字（如「使用執照」「測量成果圖」）才視為證件頁；
    長句內文提到（如「依謄本登記為主」）不算。"""
    return any(len(t) <= 14 and any(k in t for k in SENSITIVE) for t in lines)


def read_pptx(path):
    prs = Presentation(path)
    texts, cands, seen = [], [], set()
    for i, s in enumerate(prs.slides, 1):
        lines = slide_lines(s)
        if is_sensitive(lines):
            continue
        texts.append(f"[第{i}頁]\n" + "\n".join(lines))
        for shp in walk(s.shapes):
            try:
                if shp.width / 914400 < 2.5 or shp.height / 914400 < 2.0:
                    continue  # 太小：logo、圖示
                blob = shp.image.blob
                h = hashlib.md5(blob).hexdigest()
                if h in seen:
                    continue
                seen.add(h)
                im = Image.open(io.BytesIO(blob))
                if im.width * im.height > 60_000_000:
                    continue  # 超大掃描檔，不是現場照片
                im.draft("RGB", (1024, 1024))
                im = ImageOps.exif_transpose(im).convert("RGB")
                im.thumbnail((512, 512))
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=70)
                cands.append({"page": i, "blob": blob,
                              "b64": base64.b64encode(buf.getvalue()).decode()})
            except Exception:
                continue
    return "\n\n".join(texts), cands[:MAX_CAND]


def ask_gpt(text, cands):
    from openai import OpenAI
    content = [{"type": "text", "text": PROMPT + "\n\n簡報文字：\n" + text}]
    for k, c in enumerate(cands):
        content.append({"type": "text", "text": f"圖片編號 {k}（第{c['page']}頁）"})
        content.append({"type": "image_url",
                        "image_url": {"url": "data:image/jpeg;base64," + c["b64"], "detail": "low"}})
    r = OpenAI().chat.completions.create(
        model=MODEL, temperature=0,
        response_format={"type": "json_object"},
        messages=[{"role": "user", "content": content}])
    return json.loads(r.choices[0].message.content)


_D = {
    "高雄市": "楠梓區 左營區 鼓山區 三民區 鹽埕區 前金區 新興區 苓雅區 前鎮區 旗津區 小港區 鳳山區 大寮區 鳥松區 林園區 仁武區 大樹區 大社區 岡山區 路竹區 橋頭區 梓官區 彌陀區 永安區 燕巢區 田寮區 阿蓮區 茄萣區 湖內區 旗山區 美濃區 內門區 杉林區 甲仙區 六龜區 茂林區 桃源區 那瑪夏區",
    "屏東縣": "屏東市 潮州鎮 東港鎮 恆春鎮 萬丹鄉 長治鄉 麟洛鄉 九如鄉 里港鄉 鹽埔鄉 高樹鄉 萬巒鄉 內埔鄉 竹田鄉 新埤鄉 枋寮鄉 新園鄉 崁頂鄉 林邊鄉 南州鄉 佳冬鄉 琉球鄉 車城鄉 滿州鄉 枋山鄉 三地門鄉 霧臺鄉 瑪家鄉 泰武鄉 來義鄉 春日鄉 獅子鄉 牡丹鄉",
    "台南市": "新營區 鹽水區 白河區 柳營區 後壁區 東山區 麻豆區 下營區 六甲區 官田區 大內區 佳里區 學甲區 西港區 七股區 將軍區 北門區 新化區 善化區 新市區 安定區 山上區 玉井區 楠西區 南化區 左鎮區 仁德區 歸仁區 關廟區 龍崎區 永康區 安平區 安南區 中西區",
}
DIST = {d: c for c, v in _D.items() for d in v.split()}


def norm_area(a):
    """統一成「縣市＋區/鄉/鎮＋路名」。縣市以區名對照表為準（修正 AI 填錯縣市）；
    找不到區名就維持原樣，不亂改。"""
    a = re.sub(r"\s+", "", str(a or "")).replace("臺", "台")
    a = re.sub(r"[\d０-９].*$", "", a)  # 門牌、地號一律砍掉
    county = dist = rest = None
    for d in sorted(DIST, key=len, reverse=True):
        if d in a:
            county, dist, rest = DIST[d], d, a[a.rindex(d) + len(d):]
            break
    if not dist:
        m = re.match(r"(.{1,3}[縣市])(.{1,4}?[區鄉鎮市])(.*)$", a)
        if not m:
            return a
        county, dist, rest = m.groups()
    rest = re.sub(r"^[^路街巷]{1,4}?[村里]", "", rest)  # 村里不顯示
    m = re.match(r"(.*?(?:大道|路|街))((?:[一二三四五六七八九十]{1,3}段)?)", rest)
    road = (m.group(1) + m.group(2)) if m else ""
    if not road:
        m = re.match(r"(.{1,6}巷)", rest)  # 只有巷名時保留；地段（如大同段）不顯示
        road = m.group(1) if m else ""
    if len(road) > 14:
        road = ""
    return county + dist + road


def clean(d, cands):
    g = lambda k: str(d.get(k, "") or "").strip()
    if g("type") not in ("售", "租"):
        raise ValueError("無法判斷售/租：" + g("type"))
    area = norm_area(g("area"))  # 統一成 縣市＋區鄉鎮＋路名
    ids = [k for k in d.get("photo_ids", []) if isinstance(k, int) and 0 <= k < len(cands)]
    ids = list(dict.fromkeys(ids))[:12]
    return {"title": g("title"), "type": g("type"), "price": g("price"),
            "land_ping": g("land_ping"), "build_ping": g("build_ping"),
            "zoning": g("zoning"), "area": area,
            "category": g("category") if g("category") in CATS else "", "note": g("note"), "desc": g("desc")}, ids


def save_photos(pid, cands, ids):
    out = ROOT / "img" / pid
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    paths = []
    for n, k in enumerate(ids, 1):
        im = ImageOps.exif_transpose(Image.open(io.BytesIO(cands[k]["blob"]))).convert("RGB")
        im.thumbnail((1600, 1600))
        im.save(out / f"{n}.jpg", "JPEG", quality=80, optimize=True)
        paths.append(f"img/{pid}/{n}.jpg")
    return paths


def make_thumbs(data):
    """列表用的小圖 img/ID/t_檔名：缺的就補，沒有對應照片的就清掉。"""
    for l in data:
        pid = str(l.get("id", ""))
        d = ROOT / "img" / pid
        if not d.is_dir():
            continue
        keep = set()
        for p in l.get("photos") or []:
            src = ROOT / p
            if not src.is_file() or src.parent != d:
                continue
            t = d / ("t_" + src.name)
            keep.add(t.name)
            if not t.exists():
                im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
                im.thumbnail((480, 480))
                im.save(t, "JPEG", quality=72, optimize=True)
        for t in d.glob("t_*"):
            if t.name not in keep:
                t.unlink()

def cat_of(l):

    if l.get("category") in CATS:
        return l["category"]
    num = lambda k: float(re.sub(r"[^\d.]", "", str(l.get(k, ""))) or 0)
    b, d = num("build_ping"), num("land_ping")
    if not b:
        return "土地" if d else "其他"
    if not re.search("廠|倉|工業|丁種|乙種|甲種", l.get("title", "") + l.get("zoning", "")):
        return "其他"
    return "大型廠房" if b >= 500 else "小型廠房"


def jdump(o):
    return json.dumps(o, ensure_ascii=False).replace("</", "<\\/")


def price_disp(l):
    """顯示用價格：1 億以上的「XXXXX萬」換成「約X.XX億元」，其餘照原文（資料本身不改）。"""
    raw = str(l.get("price", "") or "").strip()
    m = re.match(r"^(?:總價)?\s*([\d,]+(?:\.\d+)?)\s*萬", raw)
    if not m:
        return raw
    v = float(m.group(1).replace(",", ""))
    if v < 10000:
        return raw
    yi = f"{round(v / 10000, 2):g}"
    rest = raw[m.end():].strip()
    return f"約{yi}億元" + (" " + rest if rest else "")


def loc(l):
    """回傳 (縣市簡稱, 區鄉鎮)，例如 ("高雄", "仁武區")。"""
    m = re.match(r"(.{2,3}?[市縣])(.{1,4}?[區鄉鎮市])", str(l.get("area", "")))
    if not m:
        return "", ""
    return m.group(1)[:-1], m.group(2)


def kind_of(l):
    return {"大型廠房": "廠房", "小型廠房": "廠房", "土地": "土地"}.get(cat_of(l), "物件")


def deal_of(l):
    t = l.get("title", "")
    if "租售" in t:
        return "出售出租"
    return "出租" if l.get("type") == "租" else "出售"


def seo_kw(l):
    """搜尋用短語，例如「高雄仁武區廠房出售」。"""
    c, d = loc(l)
    return f"{c}{d}{kind_of(l)}{deal_of(l)}"


def page_html(l):
    E = html.escape
    ph = [p for p in (l.get("photos") or []) if re.fullmatch(r"img/[\w\-/.]+", str(p)) and ".." not in p]
    t = l["title"] + "｜" + seo_kw(l) + "｜富住通"
    meta = [price_disp(l), l.get("area", ""), l.get("zoning", "")]
    if l.get("land_ping"): meta.append(f"土地{l['land_ping']}坪")
    if l.get("build_ping"): meta.append(f"建坪{l['build_ping']}坪")
    desc = "｜".join(x for x in meta if x) + "｜洽楊紘珉 0905-858-141"
    url = f"{SITE}/p/{l['id']}.html"
    m = re.search(r"[市縣](.+?[區鄉鎮市])", l.get("area", ""))
    ld = {"@context": "https://schema.org", "@type": "RealEstateListing", "name": t, "url": url,
          "description": desc + ("。" + l["desc"].replace("；", "，") if l.get("desc") else ""),
          "image": [f"{SITE}/{p}" for p in ph[:6]],
          "about": {"@type": "Place", "name": l.get("area", ""),
                    "address": {"@type": "PostalAddress", "addressCountry": "TW",
                                "addressLocality": m.group(1) if m else "", "streetAddress": l.get("area", "")}},
          "category": cat_of(l) + ("出租" if l.get("type") == "租" else "出售"),
          "provider": {"@type": "RealEstateAgent", "name": "富住通商用不動產 新興店 楊紘珉",
                       "telephone": "+886-905-858-141", "url": SITE}}
    rows = [("類別", cat_of(l)), ("編號", l["id"]), ("區域", l.get("area", "")), ("價格", price_disp(l)),
            ("使用分區", l.get("zoning", "")), ("土地坪數", l.get("land_ping", "")),
            ("建物坪數", l.get("build_ping", "")), ("備註", l.get("note", ""))]
    tr = "".join(f"<tr><th>{E(k)}</th><td>{E(str(v))}</td></tr>" for k, v in rows if v)
    pts = "".join(f"<li>{E(x)}</li>" for x in (l.get("desc") or "").split("；") if x)
    imgs = "".join(f'<img src="../{E(p)}" alt="{E(l["title"])}" width="400" loading="lazy">' for p in ph[:4])
    og = E(f"{SITE}/{ph[0]}") if ph else ""
    return f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(t)}</title>
<meta name="description" content="{E(desc)}">
<link rel="canonical" href="{url}">
<meta property="og:type" content="website">
<meta property="og:title" content="{E(t)}">
<meta property="og:description" content="{E(desc)}">
<meta property="og:image" content="{og}">
<meta property="og:url" content="{url}">
<meta name="twitter:card" content="summary_large_image">
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<script type="application/ld+json">{jdump(ld)}</script>
<script>location.replace("../?id={l['id']}")</script>
</head><body>
<h1>{E(l['title'])}</h1>
<p>{imgs}</p>
<table>{tr}</table>
<ul>{pts}</ul>
<p>洽詢：楊紘珉（富住通商用不動產 新興店）0905-858-141｜<a href="https://lin.ee/S6hfHqge">LINE 諮詢</a></p>
<p>相關物件：{"、".join(f'<a href="../{E(u)}">{E(lab)}</a>' for u, lab in LINKS_OF.get(l["id"], [])) or '<a href="../all.html">全部物件</a>'}</p>
<p><a href="../?id={l['id']}">查看完整物件頁</a>｜<a href="../all.html">全部物件清單</a></p>
</body></html>
"""


# ---------- 地區／類型專頁（SEO 落地頁）----------
# 只要某地區或類型有 2 筆以上「上架」物件，就自動產生一頁；少於 2 筆的頁面會自動移除。
MIN_ITEMS = 2
SLUG = {
    "仁武區": "renwu", "大社區": "dashe", "大寮區": "daliao", "岡山區": "gangshan", "路竹區": "luzhu",
    "鳳山區": "fengshan", "林園區": "linyuan", "大樹區": "dashu", "永安區": "yongan", "鳥松區": "niaosong",
    "楠梓區": "nanzi", "橋頭區": "qiaotou", "燕巢區": "yanchao", "阿蓮區": "alian", "湖內區": "hunei",
    "梓官區": "ziguan", "彌陀區": "mituo", "茄萣區": "qieding", "小港區": "xiaogang", "前鎮區": "qianzhen",
    "三民區": "sanmin", "左營區": "zuoying", "旗山區": "qishan", "美濃區": "meinong", "田寮區": "tianliao",
    "大林鎮": "dalin", "萬丹鄉": "wandan", "屏東市": "pingtung", "長治鄉": "changzhi", "麟洛鄉": "linluo",
    "內埔鄉": "neipu", "新園鄉": "xinyuan", "仁德區": "rende", "永康區": "yongkang", "安南區": "annan",
    "新市區": "xinshi", "善化區": "shanhua", "歸仁區": "guiren", "關廟區": "guanmiao",
}
ZONING_NOTE = {
    "b": "乙種工業區是都市計畫內的工業區分區，主要供公害輕微的工廠及相關設施使用，常見於市區周邊，交通與生活機能通常較方便。",
    "a": "甲種工業區是都市計畫內的工業區分區，可設置的工廠類別較廣，通常規模較大、適合製造業與重工業使用。",
    "d": "丁種建築用地是非都市土地中供工廠及相關工業設施建築使用的用地，常見於工業區或產業聚落周邊，取得面積較大的基地相對容易。",
}
TYPES = [  # (路徑, 搜尋標題, 簡短標題, 判斷函式, 說明 key)
    ("type/b-industrial.html", "高雄乙種工業區廠房・乙工用地", "乙種工業區",
     lambda l: re.search(r"乙種|乙工", l.get("zoning", "") + l.get("title", "")), "b"),
    ("type/a-industrial.html", "高雄甲種工業區廠房・甲工用地", "甲種工業區",
     lambda l: re.search(r"甲種工業|甲工", l.get("zoning", "") + l.get("title", "")), "a"),
    ("type/d-building.html", "丁種建築用地廠房・丁建廠辦", "丁種建築用地",
     lambda l: re.search(r"丁種|丁建", l.get("zoning", "") + l.get("title", "")), "d"),
    ("type/large.html", "高雄大型廠房・千坪工業用地", "大型廠房・千坪工業地",
     lambda l: cat_of(l) == "大型廠房" or _num(l.get("land_ping")) >= 1000, ""),
    ("type/rent.html", "高雄廠房出租・工業地出租", "廠房出租",
     lambda l: l.get("type") == "租" or "租" in l.get("title", ""), ""),
    ("type/land.html", "高雄工業用地・土地出售", "工業用地・土地",
     lambda l: cat_of(l) == "土地", ""),
]
LINKS_OF = {}  # 物件 id -> [(專頁路徑, 專頁短標題)]


def _num(s):
    try:
        return float(re.sub(r"[^\d.]", "", str(s or "")) or 0)
    except ValueError:
        return 0


def landing_defs(live):
    """回傳要產生的專頁清單：[{path, h1, short, items, note, kind}]"""
    pages, by_d = [], {}
    for l in live:
        c, d = loc(l)
        if d:
            by_d.setdefault((c, d), []).append(l)
    for (c, d), items in sorted(by_d.items(), key=lambda x: -len(x[1])):
        if len(items) < MIN_ITEMS:
            continue
        slug = SLUG.get(d) or "d-" + hashlib.md5(d.encode()).hexdigest()[:6]
        kinds = {kind_of(l) for l in items}
        noun = "廠房・土地" if kinds >= {"廠房", "土地"} else ("土地" if kinds == {"土地"} else "廠房")
        has_rent = any(l.get("type") == "租" or "租" in l.get("title", "") for l in items)
        pages.append({"path": f"area/{slug}.html", "h1": f"{c}{d}{noun}" + ("出售出租" if has_rent else "出售"),
                      "short": f"{d}{noun}",
                      "items": items, "note": "", "kind": "area", "place": f"{c}{d}"})
    for path, h1, short, fn, nk in TYPES:
        items = [l for l in live if fn(l)]
        if len(items) >= MIN_ITEMS:
            pages.append({"path": path, "h1": h1, "short": short, "items": items,
                          "note": ZONING_NOTE.get(nk, ""), "kind": "type", "place": ""})
    return pages


ABOUT = ('關於我們：<a href="https://fulllife5858.com.tw/" target="_blank" rel="noopener">公司官網</a>｜'
         '<a href="https://fulllife5858.com.tw/analysis.aspx" target="_blank" rel="noopener">市場分析</a>｜'
         '<a href="https://fulllife5858.com.tw/transaction.aspx" target="_blank" rel="noopener">成交指標</a>｜'
         '<a href="https://www.facebook.com/profile.php?id=61573837941258" target="_blank" rel="noopener">粉絲專頁</a>')
HEAD_NAV = """<header>
 <div class="bar"><a href="../"><img src="../logo.png" alt="富住通商用不動產 大型工業地產" width="1000" height="96"></a></div>
 <nav><a href="../">工業物件</a><a href="../price/">實價行情</a><a href="../tools/">試算工具</a><a href="../a/">廠房知識</a><a class="cta" href="../#need">幫我找廠房</a></nav>
</header>
<script>(function(){var s=location.pathname.split("/")[1]||"";document.querySelectorAll("header nav a").forEach(function(a){var h=a.getAttribute("href").replace("../","").split("/")[0];if(h&&h===s)a.className+=" on"})})()</script>"""
FOOT = """<footer><div class="in">
 <img src="../logo.png" alt="富住通商用不動產" width="1000" height="96"><br>
 <b>富茂通商用不動產股份有限公司</b>（富住通商用不動產 新興店）<br>
 營業員：楊紘珉｜(114)登字第486430號<br>
 {links}<br>
 """ + ABOUT + """<br>
 <span style="opacity:.75">本網站資料僅供參考，實際內容以現場及契約為準</span><br><span style="opacity:.75;font-size:12px">本網站使用 Google Analytics 與 Meta Pixel 蒐集匿名瀏覽統計，用於了解網站使用情形與廣告成效。</span>
</div></footer>"""
# 與 app.js 相同的 GA4 / Pixel 設定；管理者本人（is_owner）不計入
TRACK = """<script>
(function(){var OWNER=false;try{OWNER=localStorage.getItem("is_owner")==="1"}catch(e){}
window.ev=function(){};if(OWNER)return;
window.dataLayer=window.dataLayer||[];window.gtag=function(){dataLayer.push(arguments)};
var gs=document.createElement("script");gs.async=true;gs.src="https://www.googletagmanager.com/gtag/js?id=G-5YXFVRJ2JM";document.head.appendChild(gs);
gtag("js",new Date());gtag("config","G-5YXFVRJ2JM");
!function(f,b,e,v,n,t,s){if(f.fbq)return;n=f.fbq=function(){n.callMethod?n.callMethod.apply(n,arguments):n.queue.push(arguments)};if(!f._fbq)f._fbq=n;n.push=n;n.loaded=!0;n.version="2.0";n.queue=[];t=b.createElement(e);t.async=!0;t.src=v;s=b.getElementsByTagName(e)[0];s.parentNode.insertBefore(t,s)}(window,document,"script","https://connect.facebook.net/en_US/fbevents.js");
fbq("init","996673046774269");fbq("track","PageView");
var FA="";try{var lp=location.pathname;if(lp.indexOf("/a/")===0&&lp!=="/a/"&&lp!=="/a/index.html")sessionStorage.setItem("from_article",lp);FA=sessionStorage.getItem("from_article")||""}catch(e){}
window.ev=function(n){try{var p={page:location.pathname};if(FA)p.from_article=FA;gtag("event",n,p);fbq("track","Lead",{content_name:location.pathname})}catch(e){}};
document.addEventListener("click",function(e){var a=e.target.closest&&e.target.closest("a.pcard");if(a)try{gtag("event","article_click",{article:"/"+a.getAttribute("href").split("../").join(""),page:location.pathname})}catch(x){}});
})();
</script>"""


def card_html(l, n=9):
    E = html.escape
    ph = [p for p in (l.get("photos") or []) if re.fullmatch(r"img/[\w\-/.]+", str(p)) and ".." not in p]
    img = ""
    if ph:
        t = re.sub(r"[^/]+$", lambda m: "t_" + m.group(0), ph[0])
        src = t if (ROOT / t).exists() else ph[0]
        img = (f'<img class="th" src="../{E(src)}" alt="{E(l.get("title", ""))}" width="480" height="360" '
               + ('fetchpriority="high"' if n < 2 else 'loading="lazy"') + ' decoding="async">')
    meta = "｜".join(x for x in [l.get("area", ""), l.get("zoning", ""),
                                 f"土地 {l['land_ping']} 坪" if l.get("land_ping") else "",
                                 f"建坪 {l['build_ping']} 坪" if l.get("build_ping") else ""] if x)
    tag = "rent" if l.get("type") == "租" else ""
    return (f'<a class="card" href="../?id={E(l["id"])}"><div class="ph">{img or "🏭"}'
            f'<span class="tag {tag}">出{E(l.get("type", "售"))}</span></div><div class="info"><h2>{E(l["title"])}</h2>'
            f'<div class="meta"><span class="cat">{E(cat_of(l))}</span>{E(meta)}</div>'
            f'<div class="price">{E(price_disp(l))}</div></div></a>')


def range_text(items, key):
    v = sorted(_num(l.get(key)) for l in items if _num(l.get(key)))
    if not v:
        return ""
    f = lambda x: f"{x:,.0f}"
    return f"{f(v[0])} 坪" if v[0] == v[-1] else f"{f(v[0])}～{f(v[-1])} 坪"


def landing_html(pg, pages):
    E = html.escape
    items = pg["items"]
    url = f"{SITE}/{pg['path']}"
    zon = []
    for l in items:
        for z in re.split(r"[／/、,，]", re.sub(r"\(.*?\)|（.*?）", "", l.get("zoning", ""))):
            z = z.strip()
            if z and z not in zon:
                zon.append(z)
    land, build = range_text(items, "land_ping"), range_text(items, "build_ping")
    n_sale = sum(1 for l in items if l.get("type") != "租")
    n_rent = len(items) - n_sale
    where = pg["place"] or "高雄及南部"
    intro = (f"這裡整理了{where}目前上架的 {len(items)} 筆工業物件"
             f"（出售 {n_sale} 筆" + (f"、出租 {n_rent} 筆" if n_rent else "") + "）"
             + (f"，土地面積約 {land}" if land else "") + (f"，建物約 {build}" if build else "")
             + (f"，使用分區包含{'、'.join(zon[:5])}" if zon else "") + "。"
             "每筆物件都附實景照片與基本資料，點進去可以看詳細內容；"
             "如果沒有剛好符合的，也可以直接在 LINE 告訴我區域、坪數和預算，我幫您留意。")
    desc = re.sub(r"。每筆.*$", "。", intro) + "富住通商用不動產 楊紘珉 0905-858-141"
    others = [p for p in pages if p["path"] != pg["path"]]
    rel = lambda p: "../" + p["path"]
    ld = [{"@context": "https://schema.org", "@type": "ItemList", "name": pg["h1"], "url": url,
           "numberOfItems": len(items),
           "itemListElement": [{"@type": "ListItem", "position": i + 1, "url": f"{SITE}/p/{l['id']}.html",
                                "name": l["title"]} for i, l in enumerate(items)]},
          {"@context": "https://schema.org", "@type": "BreadcrumbList", "itemListElement": [
              {"@type": "ListItem", "position": 1, "name": "工業物件", "item": f"{SITE}/"},
              {"@type": "ListItem", "position": 2, "name": pg["h1"], "item": url}]}]
    links = "｜".join(f'<a href="{E(rel(p))}">{E(p["short"])}</a>' for p in others)
    return f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(pg['h1'])}｜{len(items)} 筆物件｜富住通 楊紘珉</title>
<meta name="description" content="{E(desc)}">
<link rel="canonical" href="{url}">
<meta property="og:type" content="website"><meta property="og:title" content="{E(pg['h1'])}｜富住通">
<meta property="og:description" content="{E(desc)}"><meta property="og:url" content="{url}">
{f'<meta property="og:image" content="{SITE}/{E(items[0]["photos"][0])}">' if items[0].get("photos") else ""}
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap" rel="stylesheet" media="print" onload="this.media='all'">
<link rel="stylesheet" href="../style.css?v={CSS_V}">
<script type="application/ld+json">{jdump(ld)}</script>
{TRACK}
</head><body>
{HEAD_NAV}
<main>
<a class="back" href="../">← 全部工業物件</a>
<h1>{E(pg['h1'])}</h1>
<p>{E(intro)}</p>
{f'<p class="meta">{E(pg["note"])}實際可作用途仍以主管機關核定與土地使用分區管制規定為準。</p>' if pg["note"] else ""}
<div class="grid">{"".join(card_html(l, i) for i, l in enumerate(items))}</div>
<div class="join"><div><h3>找{E(pg['short'])}？直接告訴我需求</h3><p>加入官方 LINE，告訴我區域、坪數、預算與用途，有符合的物件會第一時間通知您；新上架與降價也會通知。</p>
<div class="btns"><a class="btn line" href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener" onclick="ev('line_click')">LINE 詢問</a><a class="btn tel" href="tel:0905858141" onclick="ev('call_click')">0905-858-141</a></div></div><img class="qr" src="../line_qr.png" alt="LINE 官方帳號 QR Code"></div>
{(lambda d: f'<p>📊 <a href="../{E(PRICE_PAGES[d])}">看{E(d)}工業地・廠房實價登錄行情</a></p>' if d in PRICE_PAGES else "")(re.sub(r"^.{2}", "", pg["place"])) if pg["kind"] == "area" else ""}
{f'<h2>其他地區與類型</h2><p>{links}</p>' if links else ""}
</main>
{FOOT.format(links='<a href="../all.html">全部物件清單</a>｜<a href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener">LINE 官方帳號</a>')}
</body></html>
"""


def write_landing(live):
    pages = landing_defs(live)
    LINKS_OF.clear()
    for pg in pages:
        for l in pg["items"]:
            LINKS_OF.setdefault(l["id"], []).append((pg["path"], pg["short"]))
    keep = set()
    for pg in pages:
        f = ROOT / pg["path"]
        f.parent.mkdir(exist_ok=True)
        f.write_text(landing_html(pg, pages), "utf-8")
        keep.add(pg["path"])
    for d in ("area", "type"):
        for f in (ROOT / d).glob("*.html") if (ROOT / d).exists() else []:
            if f"{d}/{f.name}" not in keep:
                f.unlink()
    return pages


def browse_links(pages, prefix=""):
    E = html.escape
    a = [p for p in pages if p["kind"] == "area"]
    t = [p for p in pages if p["kind"] == "type"]
    out = []
    if a:
        out.append("依地區：" + "｜".join(f'<a href="{prefix}{E(p["path"])}">{E(p["short"])}</a>' for p in a))
    if t:
        out.append("依類型：" + "｜".join(f'<a href="{prefix}{E(p["path"])}">{E(p["short"])}</a>' for p in t))
    extra = []
    if PRICE_PAGES:
        extra.append(f'<a href="{prefix}price/">高雄工業地實價行情</a>')
    if ARTICLES:
        extra.append(f'<a href="{prefix}a/">廠房知識文章</a>')
    if extra:
        out.append("｜".join(extra))
    return "<br>".join(out)


def update_index(pages):
    """首頁頁尾（靜態 HTML，搜尋引擎看得到）放專頁連結。"""
    f = ROOT / "index.html"
    s = f.read_text("utf-8")
    s = re.sub(r'href="style\.css(\?v=\w+)?"', f'href="style.css?v={CSS_V}"', s)
    s = re.sub(r'src="app\.js(\?v=\w+)?"', f'src="app.js?v={JS_V}"', s)
    latest = "".join(post_card(p, ti, d, "") for p, ti, d in ARTICLES[:3])
    pblock = (f'<!--POSTS--><section class="latest" id="latest" hidden><h2>最新廠房知識</h2><div class="pcards">{latest}</div>'
              f'<p><a class="more" href="a/">看全部 {len(ARTICLES)} 篇文章 →</a></p></section><!--/POSTS-->') if ARTICLES else "<!--POSTS--><!--/POSTS-->"
    s = re.sub(r"<!--POSTS-->.*?<!--/POSTS-->", lambda _: pblock, s, flags=re.S)
    block = f"<!--BROWSE-->{browse_links(pages)}<!--/BROWSE-->"
    if "<!--BROWSE-->" in s:
        s = re.sub(r"<!--BROWSE-->.*?<!--/BROWSE-->", lambda _: block, s, flags=re.S)
        f.write_text(s, "utf-8")


# ---------- 知識文章（n8n 發完 FB 主題貼文後，寫一個 data/posts/*.json 進來）----------
POSTS = ROOT / "data" / "posts"
ARTICLES = []  # [(路徑, 標題, 日期)]，新到舊
GENERIC_HEAD = {"現場觀察", "軟性推廣", "知識分享", "產業觀察", "選址提醒", "投資觀點", "實務提醒"}


def clean_caption(cap):
    """FB 貼文 → 網站內文：拿掉標題行、LINE/電話行、hashtag 行。回傳 (標題, 段落清單, 延伸閱讀)"""
    head, paras, news = "", [], ""
    for ln in str(cap or "").splitlines():
        t = ln.strip()
        if not t:
            paras.append("")
            continue
        m = re.fullmatch(r"【(.+?)】", t)
        if m and not head:
            head = m.group(1)
            continue
        if t.startswith(("📩", "📞", "💬", "楊紘珉｜", "專注高雄工業不動產", "如需詳細照片", "---", "專業商用不動產顧問", "諮詢專線")) \
                or t in ("富住通商用不動產",) or "LINE 官方帳號" in t or re.fullmatch(r"(#\S+\s*)+", t.strip('"')):
            continue
        if "延伸閱讀" in t:
            news = re.sub(r"^\W*延伸閱讀[:：]\s*", "", t)
            continue
        if t.startswith("📰"):  # 週報的「新聞來源」行，網站改列出連結
            continue
        paras.append(t)
    blocks, cur = [], []
    for t in paras + [""]:
        if t:
            cur.append(t)
        elif cur:
            blocks.append(cur)
            cur = []
    return head, blocks, news


def _links_html(txt):
    """週報的新聞連結（格式：「1. 標題」換行「網址」）"""
    E = html.escape
    lines = [x.strip() for x in str(txt or "").splitlines() if x.strip()]
    items, title = [], ""
    for x in lines:
        if re.match(r"^https?://", x):
            if title:
                items.append(f'<li><a href="{E(x)}" target="_blank" rel="noopener nofollow">{E(title)}</a></li>')
            title = ""
        else:
            title = re.sub(r"^\d+\.\s*", "", x)
    return f'<h2>本週新聞來源</h2><ul class="alist">{"".join(items)}</ul>' if items else ""

BUL = re.compile(r"^(?:[✔✅▪•・\-]|\d+\.\s)")
SUBHEAD = re.compile(r"^(?:第[一二三四五六七八九十]+[，、]|[一二三四五六七八九十]+、|[①②③④⑤⑥⑦⑧⑨])")


def _block_html(b):
    """FB 段落 → HTML：「第一，…」「一、…」「①…」當小標；✔/・/1. 當條列"""
    E = html.escape
    out = ""
    if len(b) >= 1 and SUBHEAD.match(b[0]) and len(b[0]) <= 42:
        out += f"<h2>{E(b[0])}</h2>"
        b = b[1:]
    lead = [x for x in b if not BUL.match(x)]
    items = [x for x in b if BUL.match(x)]
    if items and all(BUL.match(x) for x in b[len(b) - len(items):]) and len(lead) == len(b) - len(items):
        if lead:
            out += "<p>" + "<br>".join(E(x) for x in lead) + "</p>"
        tag = "ol" if all(re.match(r"^\d+\.", x) for x in items) else "ul"
        out += f"<{tag}>" + "".join(f"<li>{E(re.sub(r'^(?:[✔✅▪•・-]|\d+\.)\s*', '', x))}</li>" for x in items) + f"</{tag}>"
    elif b:
        out += "<p>" + "<br>".join(E(x) for x in b) + "</p>"
    return out


def article_html(a, path, pages, rel=()):
    E = html.escape
    url = f"{SITE}/{path}"
    title = a["_title"]
    body = ""
    for b in a["_blocks"]:
        body += _block_html(b)
    desc = re.sub(r"\s+", " ", " ".join(" ".join(b) for b in a["_blocks"]))[:110]
    ld = {"@context": "https://schema.org", "@type": "Article", "headline": title, "datePublished": a["date"],
          "dateModified": a["date"], "mainEntityOfPage": url, "inLanguage": "zh-Hant",
          "author": {"@type": "Person", "name": "楊紘珉", "jobTitle": "工業不動產顧問",
                     "worksFor": {"@type": "RealEstateAgent", "name": "富住通商用不動產 新興店"}},
          "publisher": {"@type": "Organization", "name": "富住通商用不動產", "logo": {"@type": "ImageObject", "url": f"{SITE}/logo.png"}}}
    more = "｜".join(f'<a href="../{E(p["path"])}">{E(p["short"])}</a>' for p in pages)
    return f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(title)}｜廠房知識｜富住通 楊紘珉</title>
<meta name="description" content="{E(desc)}">
<link rel="canonical" href="{url}">
<meta property="og:type" content="article"><meta property="og:title" content="{E(title)}">
<meta property="og:description" content="{E(desc)}"><meta property="og:url" content="{url}">
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap" rel="stylesheet" media="print" onload="this.media='all'">
<link rel="stylesheet" href="../style.css?v={CSS_V}">
<script type="application/ld+json">{jdump(ld)}</script>
{TRACK}
</head><body>
{HEAD_NAV}
<main class="article">
<a class="back" href="./">← 廠房知識文章</a>
<h1>{E(title)}</h1>
<p class="meta">{E(a["date"])}｜楊紘珉（富住通商用不動產 工業不動產顧問）</p>
{body}
{f'<p class="meta">延伸閱讀：{E(a["_news"])}</p>' if a.get("_news") else ""}
{_links_html(a.get("links", ""))}
<div class="join"><div><h3>想找廠房、土地，或評估手上的物件？</h3><p>加入官方 LINE，直接告訴我區域、坪數、預算與用途；新物件上架也會第一時間通知您。</p>
<div class="btns"><a class="btn line" href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener" onclick="ev('line_click')">LINE 詢問</a><a class="btn tel" href="tel:0905858141" onclick="ev('call_click')">0905-858-141</a></div></div><img class="qr" src="../line_qr.png" alt="LINE 官方帳號 QR Code"></div>
{('<h2>相關文章</h2><div class="pcards">' + "".join(post_card(p, t, d, "../") for p, t, d in rel) + '</div>') if rel else ""}
{f'<h2>目前的物件</h2><p>{more}</p>' if more else ""}
</main>
{FOOT.format(links='<a href="../all.html">全部物件清單</a>｜<a href="./">廠房知識文章</a>｜<a href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener">LINE 官方帳號</a>')}
</body></html>
"""


FONT_CANDIDATES = ["/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
                   "/usr/share/fonts/opentype/noto/NotoSansCJK-Black.ttc"]


def _font(size, bold=True):
    from PIL import ImageFont
    for f in FONT_CANDIDATES if bold else [f.replace("Bold", "Regular") for f in FONT_CANDIDATES]:
        if os.path.exists(f):
            return ImageFont.truetype(f, size, index=3)  # index 3 = 繁體中文（TC）
    raise RuntimeError("找不到中文字型（GitHub Actions 需安裝 fonts-noto-cjk）")


def _wrap(draw, text, font, width):
    lines, cur = [], ""
    for ch in text:
        if draw.textlength(cur + ch, font=font) > width and cur:
            lines.append(cur)
            cur = ch.lstrip()
        else:
            cur += ch
    if cur:
        lines.append(cur)
    return lines


def render_card(a, path):
    """IG 用 1080x1350 圖卡：深藍底、標題、最多三個重點、署名與 LINE。"""
    from PIL import ImageDraw
    W, H, M = 1080, 1350, 90
    navy, red, white, soft = (30, 58, 138), (200, 48, 42), (255, 255, 255), (214, 222, 240)
    im = Image.new("RGB", (W, H), navy)
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 18], fill=red)
    strip = lambda t: re.sub(r"[\U00010000-\U0010FFFF☀-➿️]", "", t).strip()
    d.text((M, 120), "廠房知識＋", font=_font(40), fill=(255, 196, 120))
    y = 200
    tf = _font(76)
    title = strip(a["_title"]).rstrip("。")
    for ln in _wrap(d, title, tf, W - 2 * M)[:4]:
        d.text((M, y), ln, font=tf, fill=white)
        y += 104
    d.rectangle([M, y + 20, M + 120, y + 28], fill=red)
    y += 80
    pts = [strip(re.sub(r"^[✔✅▪•・\-]\s*", "", x)) for b in a["_blocks"] for x in b if re.match(r"^[✔✅▪•・\-]", x)]
    if not pts:  # 沒有條列就用第一段
        pts = [strip(" ".join(a["_blocks"][0]))] if a["_blocks"] else []
    pts = [t for t in pts if t and t.rstrip("。") != title and "你會怎麼看" not in t]
    bf = _font(42, bold=False)
    for t in pts[:3]:
        lines = _wrap(d, t, bf, W - 2 * M - 50)[:3]
        if y + 60 * len(lines) > H - 260:
            break
        d.ellipse([M, y + 18, M + 16, y + 34], fill=(255, 196, 120))
        for ln in lines:
            d.text((M + 50, y), ln, font=bf, fill=soft)
            y += 60
        y += 30
    d.rectangle([0, H - 200, W, H], fill=(22, 44, 108))
    try:
        logo = Image.open(ROOT / "logo.png").convert("RGBA")
        logo.thumbnail((620, 70))
        im.paste(logo, (M, H - 165), logo)
    except Exception:
        d.text((M, H - 165), "富住通商用不動產", font=_font(48), fill=white)
    d.text((M, H - 80), "楊紘珉｜LINE 官方帳號 @447lrpzt", font=_font(34, bold=False), fill=soft)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path, "JPEG", quality=88, optimize=True)

POST_KW = ["租", "買", "天車", "物流", "倉", "貨櫃", "挑高", "面寬", "電力", "消防", "丁種", "工業區", "土地", "擴廠", "登記", "仁武", "岡山", "橋頭", "成本"]


def _post_img(path):
    t = "a/img/t_" + path[2:-5] + ".jpg"
    return t if (ROOT / t).exists() else "a/img/" + path[2:-5] + ".jpg"


def post_card(path, title, date, prefix, eager=False):
    E = html.escape
    img = _post_img(path)
    return (f'<a class="pcard" href="{prefix}{E(path)}"><img src="{prefix}{E(img)}" alt="" {'fetchpriority="high"' if eager else 'loading="lazy"'} width="432" height="540">'
            f'<span><b>{E(title)}</b><em>{E(date)}</em></span></a>')


def _related(a, path, posts, n=3):
    """同關鍵字最多的文章優先，再補最新的"""
    txt = lambda x: x["_title"] + "".join("".join(b) for b in x["_blocks"])
    mine = {k for k in POST_KW if k in txt(a)}
    others = [(p, x) for p, x in posts if p != path]
    others.sort(key=lambda px: (len(mine & {k for k in POST_KW if k in txt(px[1])}), px[1]["date"]), reverse=True)
    return [(p, x["_title"], x["date"]) for p, x in others[:n]]


def write_articles(pages):
    ARTICLES.clear()
    out = ROOT / "a"
    posts, cards = [], set()
    for f in sorted(POSTS.glob("*.json")) if POSTS.exists() else []:
        try:
            a = json.loads(f.read_text("utf-8"))
        except Exception as e:
            print("POST FAIL", f.name, repr(e))
            continue
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(a.get("date", ""))):
            continue
        head, blocks, news = clean_caption(a.get("caption", ""))
        a["_title"] = head if head and head not in GENERIC_HEAD else (a.get("topicTitle") or head)
        a["_blocks"], a["_news"] = blocks, news
        slug = re.sub(r"[^a-z0-9-]", "", f"{a['date']}-{str(a.get('topicId', '')).lower()}").strip("-")
        cards.add(slug)
        card = ROOT / "a" / "img" / f"{slug}.jpg"
        if not card.exists():  # IG 用圖卡（範本後備的貼文也要有）
            try:
                render_card(a, card)
            except Exception as e:
                print("CARD FAIL", slug, repr(e))
        tcard = card.with_name("t_" + card.name)  # 網站列表用小圖（IG 用原圖）
        if card.exists() and not tcard.exists():
            try:
                im = Image.open(card)
                im.thumbnail((432, 540))
                im.save(tcard, "JPEG", quality=80, optimize=True)
            except Exception as e:
                print("THUMB FAIL", slug, repr(e))
        cards.add("t_" + slug)
        if a.get("source") == "template" or a.get("hidden"):
            continue  # 範本後備的短文不放網站（內容太薄）
        if sum(len("".join(b)) for b in blocks) < 80:
            continue
        posts.append((f"a/{slug}.html", a))
    if (ROOT / "a" / "img").exists():
        for f in (ROOT / "a" / "img").glob("*.jpg"):
            if f.stem not in cards:
                f.unlink()
    keep = set()
    posts.sort(key=lambda x: (x[1]["date"], str(x[1].get("topicId", ""))), reverse=True)
    ARTICLES.extend((p, a["_title"], a["date"]) for p, a in posts)
    if posts:
        out.mkdir(exist_ok=True)
    for path, a in posts:
        (ROOT / path).write_text(article_html(a, path, pages, _related(a, path, posts)), "utf-8")
        keep.add(path)
    if out.exists():
        for f in out.glob("*.html"):
            if f.name != "index.html" and f"a/{f.name}" not in keep:
                f.unlink()
    # 給首頁「最新廠房知識」與物件頁「延伸閱讀」用
    (ROOT / "data" / "posts.json").write_text(jdump([
        {"path": p, "title": a["_title"], "date": a["date"], "img": _post_img(p),
         "kw": "".join(k for k in POST_KW if k in a["_title"] + "".join("".join(b) for b in a["_blocks"])[:600])}
        for p, a in posts]), "utf-8")
    if posts:
        E = html.escape
        lis = "".join(post_card(p, t, d, "../", i < 2) for i, (p, t, d) in enumerate(ARTICLES))
        (out / "index.html").write_text(f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>廠房知識｜工業地產買賣租賃實務文章｜富住通 楊紘珉</title>
<meta name="description" content="高雄工業不動產顧問楊紘珉整理的廠房、工業用地買賣租賃實務：選址、電力、消防、使用分區與產業投資觀察。">
<link rel="canonical" href="{SITE}/a/">
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap" rel="stylesheet" media="print" onload="this.media='all'">
<link rel="stylesheet" href="../style.css?v={CSS_V}">
{TRACK}
</head><body>
{HEAD_NAV}
<main class="article"><h1>廠房知識</h1><p>買廠房、租廠房、找工業用地之前，值得先知道的實務重點。每篇都是工業不動產現場常遇到的問題。</p>
<div class="pcards">{lis}</div></main>
{FOOT.format(links='<a href="../all.html">全部物件清單</a>｜<a href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener">LINE 官方帳號</a>')}
</body></html>
""", "utf-8")
    elif (out / "index.html").exists():
        (out / "index.html").unlink()


# ---------- 實價登錄行情頁（資料由 scripts/lvr.py 從內政部開放資料下載）----------
LVR = ROOT / "data" / "lvr" / "kaohsiung_industrial.json"
PRICE_PAGES = {}  # 區名 -> 路徑
MIN_DEALS = 5


def _wan(v):
    return f"約{v / 10000:.2f}億" if v >= 10000 else f"{v:,.0f}萬"


def _zone(z):
    z = re.sub(r'^都市：其他:', '', z or '')
    return {'工': '工業區', '農': '農業區', '': '—'}.get(z, z)


def _bunit(d):
    """建坪單價＝總價÷建物坪數（含土地價值）；純土地或建坪太小時不顯示"""
    b = d.get("build_ping") or 0
    return f"{d['total_wan'] / b:.1f}".rstrip("0").rstrip(".") if d["kind"] != "土地" and b >= 5 else "—"


def _bp(d):
    return f"{d['build_ping']:,.0f}" if d.get('build_ping') else '—'


def _deal_unit(d):
    return round(d["total_wan"] / d["land_ping"], 1) if d.get("land_ping") else 0


def _median(v):
    v = sorted(v)
    n = len(v)
    return 0 if not n else (v[n // 2] if n % 2 else round((v[n // 2 - 1] + v[n // 2]) / 2, 1))


def _stats(deals, since):
    rec = [d for d in deals if not _special(d) and d["date"] >= since]
    land = [_deal_unit(d) for d in rec if d["kind"] == "土地" and _deal_unit(d)]
    bld = [_deal_unit(d) for d in rec if d["kind"] != "土地" and _deal_unit(d)]
    return len(rec), (_median(land), len(land)), (_median(bld), len(bld))


def _stat_text(v, n):
    if not n:
        return "近一年無成交"
    return f"{v:g} 萬／地坪（{n} 筆）" + ("，筆數少僅供參考" if n < 3 else "")


# ---------- 實價登錄備註解析：特殊交易、土地建物分開計價 ----------
HARD = [  # 會讓價格失真的交易，不列入中位數與圖表（表格仍列出並標示）
    ("親友／特殊關係", r"親友|特殊關係|關係人|二親等"),
    ("政府標售／承購", r"政府機關標讓售|協議價購|水利地承購|承租轉承購|向政府機關"),
    ("含公設保留地", r"公共設施保留地"),
    ("畸零地", r"畸零地"),
    ("地上權", r"地上權"),
    ("含租約", r"含租約|帶租約"),
    ("土地建物分件登記", r"分件登記"),
    ("其他特殊", r"債權債務|急買急賣|瑕疵|凶宅|無償贈與"),
]
SOFT = [("含增建", r"增建|未登記建物|夾層|頂樓加蓋|陽台外推"), ("含設備", r"傢俱|設備|機電")]


def _amt(note, label):
    """從備註抓「土地：1,234萬5,678元」這類金額，回傳萬元；抓不到回傳 None"""
    m = re.search(label + r"[^0-9]{0,12}?([\d,]+(?:\.\d+)?)\s*(萬)?\s*(?:([\d,]+)\s*元)?", note)
    if not m:
        return None
    a = float(m.group(1).replace(",", ""))
    if m.group(2):
        return a + (float(m.group(3).replace(",", "")) / 10000 if m.group(3) else 0)
    return a / 10000 if a >= 100000 else None  # 只有數字沒有「萬」，視為元


def note_info(d):
    if "_ni" in d:
        return d["_ni"]
    n = d.get("note", "") or ""
    hard = [t for t, r in HARD if re.search(r, n)]
    soft = [t for t, r in SOFT if re.search(r, n)]
    land_w, bld_w = _amt(n, r"土地"), _amt(n, r"(?:建物|房屋)")
    split = None
    if land_w and bld_w and d.get("total_wan") and abs(land_w + bld_w - d["total_wan"]) <= d["total_wan"] * 0.15:
        split = {"land": round(land_w / d["land_ping"], 1) if d.get("land_ping") else None,
                 "bld": round(bld_w / d["build_ping"], 1) if d.get("build_ping") else None}
    d["_ni"] = {"hard": hard, "soft": soft, "split": split, "special": bool(hard) or (d.get("special") and not soft)}
    return d["_ni"]


def _special(d):
    return note_info(d)["special"]


def _tags_html(d):
    E = html.escape
    ni = note_info(d)
    out = [f'<span class="tg sp">{E(t)}</span>' for t in ni["hard"]]
    if ni["special"] and not ni["hard"]:
        out.append('<span class="tg sp">特殊交易</span>')
    if ni["split"]:
        parts = []
        if ni["split"]["land"]:
            parts.append(f'土地 {ni["split"]["land"]:g}')
        if ni["split"]["bld"]:
            parts.append(f'建物 {ni["split"]["bld"]:g}')
        out.append(f'<span class="tg sv">分開計價：{"、".join(parts)} 萬／坪</span>')
    out += [f'<span class="tg">{E(t)}</span>' for t in ni["soft"]]
    return "".join(out) or "—"


def price_page_html(dist, deals, upd, pages, listings):
    E = html.escape
    path = PRICE_PAGES[dist]
    url = f"{SITE}/{path}"
    since = (datetime.date.fromisoformat(upd) - datetime.timedelta(days=365)).isoformat()
    n, (lm, ln), (bm, bn) = _stats(deals, since)
    ok = [d for d in deals if not _special(d)]
    sp = len(deals) - len(ok)
    rows = "".join(
        f"<tr{' class=\"spr\"' if _special(d) else ''}><td>{E(d['date'][:7])}</td><td>{E(d['road'] or '—')}</td><td>{'土地' if d['kind'] == '土地' else '房地'}</td>"
        f"<td>{E(_zone(d['zone']))}</td><td>{d['land_ping']:,.0f}</td>"
        f"<td>{_bp(d)}</td>"
        f"<td>{E(_wan(d['total_wan']))}</td><td><b>{(f"{_deal_unit(d):g}" if _deal_unit(d) else "—")}</b></td><td>{_bunit(d)}</td><td class=\"tgs\">{_tags_html(d)}</td></tr>"
        for d in deals[:50])
    area = next((p for p in pages if p["kind"] == "area" and p["place"].endswith(dist)), None)
    mine = [l for l in listings if loc(l)[1] == dist]
    mine_html = ""
    if area:
        mine_html = f'<p>👉 目前我在{E(dist)}的物件：<a href="../{E(area["path"])}">{E(area["h1"])}</a></p>'
    elif mine:
        mine_html = "<p>👉 目前我在" + E(dist) + "的物件：" + "、".join(
            f'<a href="../p/{E(l["id"])}.html">{E(l["title"])}</a>' for l in mine) + "</p>"
    others = "｜".join(f'<a href="../{E(p)}">{E(k)}</a>' for k, p in PRICE_PAGES.items() if k != dist)
    title = f"高雄{dist}工業地・廠房實價登錄行情"
    desc = (f"{dist}近一年工業類成交 {n} 筆。土地：{_stat_text(lm, ln)}；廠房（房地）：{_stat_text(bm, bn)}。"
            f"資料來源內政部實價登錄，{upd} 更新。")
    ld = [{"@context": "https://schema.org", "@type": "WebPage", "name": title, "url": url, "dateModified": upd,
           "description": desc},
          {"@context": "https://schema.org", "@type": "BreadcrumbList", "itemListElement": [
              {"@type": "ListItem", "position": 1, "name": "工業物件", "item": f"{SITE}/"},
              {"@type": "ListItem", "position": 2, "name": "實價登錄行情", "item": f"{SITE}/price/"},
              {"@type": "ListItem", "position": 3, "name": title, "item": url}]}]
    return f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(title)}｜{E(upd[:7])} 更新｜富住通 楊紘珉</title>
<meta name="description" content="{E(desc)}">
<link rel="canonical" href="{url}">
<meta property="og:type" content="website"><meta property="og:title" content="{E(title)}"><meta property="og:description" content="{E(desc)}"><meta property="og:url" content="{url}">
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap" rel="stylesheet" media="print" onload="this.media='all'">
<link rel="stylesheet" href="../style.css?v={CSS_V}">
<script type="application/ld+json">{jdump(ld)}</script>
{TRACK}
</head><body>
{HEAD_NAV}
<main>
<a class="back" href="./">← 高雄工業地產實價行情</a>
<h1>{E(title)}</h1>
<div class="stats"><div><span>近一年成交</span><b>{n} 筆</b></div><div><span>土地 地坪單價中位數</span><b>{E(_stat_text(lm, ln))}</b></div><div><span>廠房（房地） 地坪單價中位數</span><b>{E(_stat_text(bm, bn))}</b></div></div>
<p class="meta">資料來源：內政部實價登錄開放資料，{E(upd)} 更新。地坪單價＝總價÷土地坪數，建坪單價＝總價÷建物坪數（兩者都含土地與建物價值，廠房小、土地大時建坪單價會偏高），單位：萬元。灰色列是特殊交易（親友、政府標售、含公設保留地、含租約等，依實價登錄備註判斷），共 {sp} 筆，不列入中位數與圖表；備註寫明土地、建物分開計價的，另外標出各自的每坪單價。位置只顯示到路名，土地交易不顯示位置。</p>
{mine_html}
{trend_svg(deals)}
{scatter_svg(deals)}
<h2>近期成交明細</h2>
<div class="tbl"><table class="deals"><thead><tr><th>年月</th><th>位置</th><th>標的</th><th>分區</th><th>土地坪</th><th>建坪</th><th>總價</th><th>地坪單價</th><th>建坪單價</th><th>備註</th></tr></thead><tbody>{rows}</tbody></table></div>
<div class="join"><div><h3>想知道你的廠房、土地現在值多少？</h3><p>實價登錄只看得到成交價，看不到屋況、面寬、電力和路寬。加 LINE 告訴我地段與坪數，我幫你對照近期成交，免費給你行情建議。</p>
<div class="btns"><a class="btn line" href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener" onclick="ev('line_click')">LINE 免費估價</a><a class="btn tel" href="tel:0905858141" onclick="ev('call_click')">0905-858-141</a></div></div><img class="qr" src="../line_qr.png" alt="LINE 官方帳號 QR Code"></div>
{f'<h2>其他地區實價行情</h2><p>{others}</p>' if others else ""}
</main>
{FOOT.format(links='<a href="../all.html">全部物件清單</a>｜<a href="./">實價行情總覽</a>｜<a href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener">LINE 官方帳號</a>')}
</body></html>
"""


# ---------- 行情圖表（純 SVG，不需外部套件）----------
C_BLD, C_LAND = "#2a78d6", "#d97706"  # 房地／土地（已用 dataviz 驗證色盲可分辨）


def _half(d):
    y, m = int(d[:4]), int(d[5:7])
    return f"{y}{'上' if m <= 6 else '下'}"


def trend_svg(deals):
    ok = [d for d in deals if not _special(d) and _deal_unit(d)]
    if not ok:
        return ""
    b = {}
    for d in ok:
        b.setdefault(_half(d["date"]), []).append(_deal_unit(d))
    keys = sorted(b, key=lambda k: (k[:4], k[4] == "下"))
    pts = [(k, _median(b[k]), len(b[k])) for k in keys]
    if len(pts) < 3:
        return ""
    W, H, L, R, T, B = 640, 260, 52, 70, 20, 40
    vmax = max(v for _, v, _ in pts) * 1.15
    step = 10 if vmax <= 60 else 20 if vmax <= 120 else 50
    vmax = (int(vmax // step) + 1) * step
    x = lambda i: L + (W - L - R) * (i / (len(pts) - 1))
    y = lambda v: T + (H - T - B) * (1 - v / vmax)
    g = "".join(f'<line x1="{L}" x2="{W-R}" y1="{y(t):.1f}" y2="{y(t):.1f}" class="gr"/><text x="{L-8}" y="{y(t)+4:.1f}" class="ax" text-anchor="end">{t}</text>'
                for t in range(0, vmax + 1, step))
    xl = "".join(f'<text x="{x(i):.1f}" y="{H-14}" class="ax" text-anchor="middle">{k[:4]}{k[4]}</text>' for i, (k, _, _) in enumerate(pts))
    path = " ".join(f"{'M' if i == 0 else 'L'}{x(i):.1f},{y(v):.1f}" for i, (_, v, _) in enumerate(pts))
    dots = "".join(f'<g class="pt"><circle cx="{x(i):.1f}" cy="{y(v):.1f}" r="12" fill="transparent"/><circle cx="{x(i):.1f}" cy="{y(v):.1f}" r="4.5" fill="{C_BLD}" stroke="#fff" stroke-width="2"/>'
                   f'<title>{k[:4]} {"上半年" if k[4] == "上" else "下半年"}：中位數 {v:g} 萬／地坪（{n} 筆）</title></g>'
                   for i, (k, v, n) in enumerate(pts))
    last = pts[-1]
    lab = f'<text x="{x(len(pts)-1)+10:.1f}" y="{y(last[1])+5:.1f}" class="lb">{last[1]:g} 萬</text>'
    return (f'<figure class="chart"><figcaption>地坪單價中位數走勢（半年，萬／地坪）</figcaption>'
            f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="地坪單價中位數走勢">{g}{xl}'
            f'<path d="{path}" fill="none" stroke="{C_BLD}" stroke-width="2"/>{dots}{lab}</svg>'
            f'<p class="meta">每個點是該半年所有成交（土地＋房地）的中位數，滑過或點一下看筆數。</p></figure>')


def scatter_svg(deals):
    import math
    ok = [d for d in deals if not _special(d) and _deal_unit(d) >= 3 and d.get("land_ping", 0) >= 10]  # 單價過低多為持分、道路用地
    if len(ok) < 5:
        return ""
    W, H, L, R, T, B = 640, 300, 52, 16, 20, 44
    xs = [d["land_ping"] for d in ok]
    us = sorted(_deal_unit(d) for d in ok)
    p95 = us[int(len(us) * 0.95) - 1] if len(us) > 1 else us[0]
    vmax = max(p95 * 1.3, 10)
    step = 10 if vmax <= 60 else 20 if vmax <= 120 else 50
    vmax = (int(vmax // step) + 1) * step
    xs_s = sorted(xs)
    xcap = xs_s[min(len(xs_s) - 1, int(len(xs_s) * 0.98))]
    ok = [d for d in ok if d["land_ping"] <= xcap]
    lx0, lx1 = math.log10(max(min(xs), 10)), math.log10(xcap * 1.15)
    if lx1 - lx0 < 1:
        lx1 = lx0 + 1
    x = lambda v: L + (W - L - R) * ((math.log10(v) - lx0) / (lx1 - lx0))
    y = lambda v: T + (H - T - B) * (1 - min(v, vmax) / vmax)
    g = "".join(f'<line x1="{L}" x2="{W-R}" y1="{y(t):.1f}" y2="{y(t):.1f}" class="gr"/><text x="{L-8}" y="{y(t)+4:.1f}" class="ax" text-anchor="end">{t}</text>'
                for t in range(0, vmax + 1, step))
    ticks = [t for t in (10, 30, 100, 300, 1000, 3000, 10000) if lx0 <= math.log10(t) <= lx1]
    xl = "".join(f'<text x="{x(t):.1f}" y="{H-22}" class="ax" text-anchor="middle">{t:,}</text>' for t in ticks)
    xl += f'<text x="{(L+W-R)/2:.0f}" y="{H-4}" class="ax" text-anchor="middle">土地坪數（對數刻度）</text>'
    dots = ""
    for d in sorted(ok, key=lambda d: d["date"]):
        u = _deal_unit(d)
        c = C_LAND if d["kind"] == "土地" else C_BLD
        k = "土地" if d["kind"] == "土地" else "房地"
        over = u > vmax
        dots += (f'<g class="pt"><circle cx="{x(d["land_ping"]):.1f}" cy="{y(u):.1f}" r="5" fill="{c}" fill-opacity=".85" stroke="#fff" stroke-width="1.5"{" stroke-dasharray=\"2 2\"" if over else ""}/>'
                 f'<title>{d["date"][:7]}｜{d["road"] or "（未揭露位置）"}｜{k}｜土地 {d["land_ping"]:,.0f} 坪｜{_wan(d["total_wan"])}｜{u:g} 萬／地坪{"（超出圖表上緣）" if over else ""}</title></g>')
    leg = (f'<div class="legend"><span><i style="background:{C_BLD}"></i>房地（廠房＋土地）</span>'
           f'<span><i style="background:{C_LAND}"></i>純土地</span></div>')
    return (f'<figure class="chart"><figcaption>每筆成交：土地坪數 × 地坪單價（萬／地坪）</figcaption>{leg}'
            f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="土地坪數與地坪單價散佈圖">{g}{xl}{dots}</svg>'
            f'<p class="meta">對照自己的坪數，看同區類似大小的成交落在哪裡。滑過或點一下圓點看明細。</p></figure>')


def map_block(by, upd, live):
    """總覽頁的區行情地圖（Leaflet + 開放地圖），資料寫在 data/lvr/map.json"""
    since = (datetime.date.fromisoformat(upd) - datetime.timedelta(days=365)).isoformat()
    m = {}
    for k, v in by.items():
        rec = [d for d in v if not _special(d) and d["date"] >= since and _deal_unit(d)]
        n, (lm, ln), (bm, bn) = _stats(v, since)
        m[k] = {"n": n, "all": _median([_deal_unit(d) for d in rec]) if len(rec) >= 3 else None,
                "land": lm if ln else None, "ln": ln, "bld": bm if bn else None, "bn": bn,
                "page": PRICE_PAGES.get(k, ""), "mine": []}
    for l in live:
        c, d = loc(l)
        if c == "高雄" and d in m:
            m[d]["mine"].append({"id": l["id"], "t": l["title"]})
        elif c == "高雄" and d:
            m[d] = {"n": 0, "all": None, "land": None, "ln": 0, "bld": None, "bn": 0, "page": "", "mine": [{"id": l["id"], "t": l["title"]}]}
    (ROOT / "data" / "lvr" / "map.json").write_text(json.dumps(m, ensure_ascii=False), "utf-8")
    return """<div class="mapwrap"><div id="pmap" role="img" aria-label="高雄各區工業地地坪單價地圖"></div><div class="maplegend" id="pleg"></div></div>
<p class="meta">顏色越深＝近一年工業地・廠房地坪單價中位數越高（萬／地坪）；灰色＝成交少於 3 筆。🏭＝目前有我的物件。點地區看詳細。</p>
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.css">
<script src="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.js"></script>
<script>
(function(){
var RAMP=["#b7d3f6","#86b6ef","#3987e5","#1c5cab","#0d366b"],NODATA="#e5e4df";
function esc(s){return String(s==null?"":s).replace(/[&<>"']/g,function(c){return {"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]})}
Promise.all([fetch("../data/lvr/kaohsiung_districts.geojson").then(function(r){return r.json()}),fetch("../data/lvr/map.json").then(function(r){return r.json()})]).then(function(a){
 var geo=a[0],M=a[1],vals=Object.keys(M).map(function(k){return M[k].all}).filter(function(v){return v}).sort(function(x,y){return x-y});
 var br=[1,2,3,4].map(function(i){return vals[Math.floor(vals.length*i/5)]});
 function col(v){if(!v)return NODATA;for(var i=0;i<4;i++)if(v<br[i])return RAMP[i];return RAMP[4]}
 var map=L.map("pmap",{scrollWheelZoom:false,attributionControl:true,zoomSnap:0.25});
 L.tileLayer("https://wmts.nlsc.gov.tw/wmts/EMAP/default/GoogleMapsCompatible/{z}/{y}/{x}",{maxZoom:17,opacity:.55,attribution:'底圖：<a href="https://maps.nlsc.gov.tw/" target="_blank" rel="noopener">內政部國土測繪中心</a>｜行政區界：內政部'}).addTo(map);
 var focus=[];
 var lay=L.geoJSON(geo,{style:function(f){var d=M[f.properties.name]||{};return {color:"#ffffff",weight:1.5,fillColor:col(d.all),fillOpacity:d.all?.78:.45}},
  onEachFeature:function(f,l){var k=f.properties.name,d=M[k]||{};if(d.n||(d.mine&&d.mine.length))focus.push(l);
   var h="<b>"+esc(k)+"</b><br>近一年成交："+(d.n||0)+" 筆";
   if(d.all)h+="<br>地坪單價中位數：<b>"+d.all+" 萬</b>";
   if(d.land)h+="<br>土地："+d.land+" 萬（"+d.ln+" 筆）";if(d.bld)h+="<br>廠房（房地）："+d.bld+" 萬（"+d.bn+" 筆）";
   if(d.page)h+='<br><a href="../'+esc(d.page)+'">看'+esc(k)+'每一筆成交 →</a>';
   if(d.mine&&d.mine.length)h+="<br>🏭 我的物件："+d.mine.map(function(x){return '<a href="../p/'+esc(x.id)+'.html">'+esc(x.t)+"</a>"}).join("、");
   l.bindPopup(h);l.bindTooltip(esc(k)+(d.all?"｜"+d.all+" 萬":""),{sticky:true});
   l.on("mouseover",function(){l.setStyle({weight:3,color:"#1e3a8a"})});l.on("mouseout",function(){lay.resetStyle(l)});
   if(d.mine&&d.mine.length){var c=l.getBounds().getCenter();L.marker(c,{icon:L.divIcon({className:"mine",html:"🏭<b>"+d.mine.length+"</b>",iconSize:[38,22]})}).addTo(map).bindPopup(h)}
  }}).addTo(map);
 var fb=L.featureGroup(focus.length?focus:[lay]).getBounds();map.fitBounds(fb,{padding:[6,6]});
 var lg="";var lo=[vals[0]].concat(br);for(var i=0;i<5;i++){lg+='<span><i style="background:'+RAMP[i]+'"></i>'+(i<4?Math.round(lo[i])+"–"+Math.round(br[i]):Math.round(br[3])+"+")+"</span>"}
 lg+='<span><i style="background:'+NODATA+'"></i>筆數不足</span>';document.getElementById("pleg").innerHTML="萬／地坪："+lg;
}).catch(function(){document.getElementById("pmap").style.display="none"});
})();
</script>"""


def write_price_pages(pages, live):
    PRICE_PAGES.clear()
    out = ROOT / "price"
    if not LVR.exists():
        return
    data = json.loads(LVR.read_text("utf-8"))
    upd = data.get("updated") or datetime.date.today().isoformat()
    by = {}
    for d in data.get("deals", []):
        by.setdefault(d["dist"], []).append(d)
    dists = [k for k, v in sorted(by.items(), key=lambda x: -len(x[1]))
             if k and sum(1 for d in v if not _special(d)) >= MIN_DEALS]
    for k in dists:
        PRICE_PAGES[k] = f"price/{SLUG.get(k) or 'd-' + hashlib.md5(k.encode()).hexdigest()[:6]}.html"
    out.mkdir(exist_ok=True)
    keep = set()
    for k in dists:
        (ROOT / PRICE_PAGES[k]).write_text(price_page_html(k, by[k], upd, pages, live), "utf-8")
        keep.add(PRICE_PAGES[k])
    for f in out.glob("*.html"):
        if f.name != "index.html" and f"price/{f.name}" not in keep:
            f.unlink()
    map_html = map_block(by, upd, live)
    (ROOT / "data" / "lvr" / "pages.json").write_text(json.dumps(PRICE_PAGES, ensure_ascii=False), "utf-8")  # 給 app.js 用
    E = html.escape
    since = (datetime.date.fromisoformat(upd) - datetime.timedelta(days=365)).isoformat()
    trs = ""
    for k in dists:
        n, (lm, ln), (bm, bn) = _stats(by[k], since)
        trs += (f'<tr><td><a href="../{E(PRICE_PAGES[k])}">{E(k)}</a></td><td>{n}</td>'
                f'<td>{f"{lm:g}" if ln else "—"}</td><td>{f"{bm:g}" if bn else "—"}</td></tr>')
    (out / "index.html").write_text(f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>高雄工業地產實價登錄行情｜各區工業地・廠房成交價｜{E(upd[:7])} 更新</title>
<meta name="description" content="高雄各區工業區、丁種建築用地、廠房的實價登錄成交行情，依地區整理地坪單價中位數與近期成交。資料來源內政部實價登錄，{E(upd)} 更新。">
<link rel="canonical" href="{SITE}/price/">
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap" rel="stylesheet" media="print" onload="this.media='all'">
<link rel="stylesheet" href="../style.css?v={CSS_V}">
{TRACK}
</head><body>
{HEAD_NAV}
<main><h1>高雄工業地產實價登錄行情</h1>
<p>整理高雄各區工業區、丁種建築用地與廠房的實價登錄成交，每 10 天自動更新。點地區看近期每一筆成交。</p>
{map_html}
<h2>各區行情一覽</h2>
<div class="tbl"><table class="deals"><thead><tr><th>地區</th><th>近一年成交筆數</th><th>土地中位數<br>（萬／地坪）</th><th>廠房中位數<br>（萬／地坪）</th></tr></thead><tbody>{trs}</tbody></table></div>
<p class="meta">資料來源：內政部實價登錄開放資料，{E(upd)} 更新。地坪單價＝總價÷土地坪數（房地含建物價值），已排除政府標售、親友等特殊交易。</p>
<div class="join"><div><h3>想知道你的廠房、土地現在值多少？</h3><p>加 LINE 告訴我地段與坪數，我幫你對照近期成交，免費給你行情建議。</p>
<div class="btns"><a class="btn line" href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener" onclick="ev('line_click')">LINE 免費估價</a></div></div><img class="qr" src="../line_qr.png" alt="LINE 官方帳號 QR Code"></div>
</main>
{FOOT.format(links='<a href="../all.html">全部物件清單</a>｜<a href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener">LINE 官方帳號</a>')}
</body></html>
""", "utf-8")


def write_share_pages(data):
    pdir = ROOT / "p"
    pdir.mkdir(exist_ok=True)
    pub = [l for l in data
           if l.get("status") != "待確認" and re.fullmatch(r"[A-Za-z0-9_-]{1,40}", str(l.get("id", "")))]
    if not data:  # 資料為空時不動（防止 listings.json 被清空時誤刪）
        return
    live_l = [l for l in pub if l.get("status") == "上架"]
    write_price_pages([], live_l)  # 先算出行情頁路徑，地區專頁才能連過去
    pages = write_landing(live_l)
    write_articles(pages)
    write_price_pages(pages, live_l)
    live = set()
    for l in pub:
        live.add(f"{l['id']}.html")
        (pdir / f"{l['id']}.html").write_text(page_html(l), "utf-8")
    # 清掉不再公開的舊分享頁
    for f in pdir.glob("*.html"):
        if f.name not in live:
            f.unlink()
    update_index(pages)
    write_site_files(pub, pages)


def write_tools_page():
    body = ROOT / "scripts" / "tools_body.html"
    if not body.exists():
        return False
    (ROOT / "tools").mkdir(exist_ok=True)
    desc = "廠房租金與稅費拆算（個人／法人房東、10% 扣繳、二代健保、5% 營業稅）、置產投報率、建蔽率容積率試算。結果可複製或分享給老闆。富住通商用不動產 楊紘珉。"
    ld = {"@context": "https://schema.org", "@type": "WebApplication", "name": "工業地產試算工具", "url": f"{SITE}/tools/",
          "applicationCategory": "FinanceApplication", "operatingSystem": "Any", "inLanguage": "zh-Hant",
          "offers": {"@type": "Offer", "price": "0", "priceCurrency": "TWD"}, "description": desc}
    (ROOT / "tools" / "index.html").write_text(f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>廠房租金稅費試算・投報率・建蔽率容積率計算機｜富住通 楊紘珉</title>
<meta name="description" content="{html.escape(desc)}">
<link rel="canonical" href="{SITE}/tools/">
<meta property="og:type" content="website"><meta property="og:title" content="工業地產試算工具｜租金稅費・投報率・建蔽容積"><meta property="og:description" content="{html.escape(desc)}"><meta property="og:url" content="{SITE}/tools/">
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap" rel="stylesheet" media="print" onload="this.media='all'">
<link rel="stylesheet" href="../style.css?v={CSS_V}">
<script type="application/ld+json">{jdump(ld)}</script>
{TRACK}
</head><body>
{HEAD_NAV}
{body.read_text("utf-8")}
{FOOT.format(links='<a href="../all.html">全部物件清單</a>｜<a href="../price/">實價行情</a>｜<a href="https://lin.ee/S6hfHqge" target="_blank" rel="noopener">LINE 官方帳號</a>')}
</body></html>
""", "utf-8")
    return True


def write_site_files(pub, pages=()):
    E = html.escape
    live = [l for l in pub if l.get("status") == "上架"]
    sec = ""
    for c in CATS:
        items = [l for l in live if cat_of(l) == c]
        if items:
            sec += f"<h2>{c}</h2><ul>" + "".join(
                f'<li><a href="p/{l["id"]}.html">{E(l["title"])}</a>｜{E(l.get("area",""))}｜{E(price_disp(l))}</li>'
                for l in items) + "</ul>"
    if pages:
        sec += "<h2>依地區・類型瀏覽</h2><p>" + browse_links(pages) + "</p>"
    (ROOT / "all.html").write_text(f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>高雄工業廠房・工業用地出售出租物件總覽｜富住通</title>
<link rel="icon" href="/favicon.png" type="image/png"><link rel="apple-touch-icon" href="/apple-touch-icon.png">
<meta name="description" content="富住通商用不動產 新興店 楊紘珉，高雄、仁武、大寮、岡山等地工業廠房與工業用地出售、出租物件清單，含土地坪數、建坪與價格，LINE 即時諮詢。">
<link rel="canonical" href="{SITE}/all.html">
</head><body><h1>工業廠房・工業用地物件總覽</h1>{sec or "<p>目前沒有上架物件</p>"}
<p><a href="./">回首頁</a>｜洽詢：楊紘珉 0905-858-141｜<a href="https://lin.ee/S6hfHqge">LINE 諮詢</a></p></body></html>
""", "utf-8")
    urls = ([f"{SITE}/", f"{SITE}/all.html"] + [f"{SITE}/{p['path']}" for p in pages]
            + [f"{SITE}/p/{l['id']}.html" for l in live]
            + ([f"{SITE}/a/"] + [f"{SITE}/{p}" for p, _, _ in ARTICLES] if ARTICLES else [])
            + ([f"{SITE}/price/"] + [f"{SITE}/{p}" for p in PRICE_PAGES.values()] if PRICE_PAGES else [])
            + ([f"{SITE}/tools/"] if write_tools_page() else []))
    (ROOT / "sitemap.xml").write_text('<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        + "".join(f"<url><loc>{E(u)}</loc></url>\n" for u in urls) + "</urlset>\n", "utf-8")
    (ROOT / "robots.txt").write_text(f"User-agent: *\nAllow: /\nDisallow: /admin.html\n\nSitemap: {SITE}/sitemap.xml\n", "utf-8")


def main():
    old_f = ROOT / "listings.json"  # 舊位置：搬遷前相容
    f = DATA
    f.parent.mkdir(exist_ok=True)
    src_f = f if f.exists() else old_f
    data = json.loads(src_f.read_text("utf-8")) if src_f.exists() else []
    data = [l for l in data if not l["title"].startswith("【範例】")]
    failed = 0
    for p in sorted((ROOT / "pptx").glob("*")):
        if p.suffix.lower() != ".pptx":
            continue
        try:
            src = re.sub(r"^\d{14}__", "", p.name)  # n8n 上傳時會加時間前綴避免同名衝突
            text, cands = read_pptx(p)
            d, ids = clean(ask_gpt(text, cands), cands)
            old = next((l for l in data if l.get("src") == src), None)
            if old:
                pid = old["id"]
            else:
                nums = [int(l["id"][4:]) for l in data if l["id"][4:].isdigit()]
                pid = "IND-%03d" % (max(nums, default=0) + 1)
            rec = {"id": pid, "src": src, **d,
                   "status": old["status"] if old else "待確認",
                   "photos": save_photos(pid, cands, ids)}
            if old:
                data[data.index(old)] = rec
            else:
                data.append(rec)
            print("OK", pid, p.name, len(ids), "photos")
            p.unlink()  # 成功才刪除；失敗會保留檔案方便重試
        except Exception as e:
            failed += 1
            print("FAIL", p.name, repr(e))
    for l in data:  # 統一地址格式（已符合的不會變動）
        n = norm_area(l.get("area", ""))
        if n and n != l.get("area"):
            print("AREA", l["id"], l.get("area"), "->", n)
            l["area"] = n
    f.write_text(json.dumps(data, ensure_ascii=False, indent=2), "utf-8")
    if not f.exists() or old_f.exists():
        old_f.unlink(missing_ok=True)
    make_thumbs(data)
    write_share_pages(data)
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
